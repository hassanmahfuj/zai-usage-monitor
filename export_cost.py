"""Generate monthly cost allocation Excel report for Z.ai usage data."""

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from formatting import TOKEN_NUMFMT


def _style_header(ws, row: int, ncols: int) -> None:
    """Apply header styling to a row."""
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border


def _style_data_rows(ws, start_row: int, end_row: int, ncols: int, left_cols: int = 1) -> None:
    """Apply alternating row styling to data rows.

    Args:
        left_cols: Number of leading columns to left-align (e.g. S.No + User).
                   Remaining columns are right-aligned (numeric).
    """
    light_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for row in range(start_row, end_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if col <= left_cols else "right")
            if (row - start_row) % 2 == 1:
                cell.fill = light_fill


def _auto_width(ws, ncols: int) -> None:
    """Auto-fit column widths."""
    for col in range(1, ncols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 30)


def generate_cost_excel(
    df: pd.DataFrame,
    start_date: date,
    end_date: date,
    cdf: pd.DataFrame | None = None,
) -> bytes:
    """Generate a cost allocation Excel workbook.

    Creates a workbook with 4-5 sheets:
      1. Monthly Summary - total cost, tokens, requests (and contributions) per month
      2. Per-User Allocation - who used what, for finance (and contributions)
      3. Per-Model Breakdown - which models cost the most
      4. Raw Data - all filtered records
      5. Contributions - daily per-user contribution counts (only when cdf provided)

    Args:
        df: Filtered usage DataFrame.
        start_date: Report period start date.
        end_date: Report period end date.
        cdf: Optional filtered contributions DataFrame with columns
            username, contribution_date, count.

    Returns:
        Excel file content as bytes.
    """
    if cdf is None:
        cdf = pd.DataFrame()
    have_contrib = not cdf.empty

    wb = Workbook()

    # --- Sheet 1: Monthly Summary ---
    ws1 = wb.active
    ws1.title = "Monthly Summary"

    df_copy = df.copy()
    df_copy["month"] = pd.to_datetime(df_copy["billing_date"]).dt.to_period("M").astype(str)

    monthly = (
        df_copy.groupby("month", as_index=False)
        .agg(tokens=("token_usage", "sum"), cost=("cost", "sum"), requests=("requests", "sum"))
        .sort_values("month")
    )

    if have_contrib:
        cdf_copy = cdf.copy()
        cdf_copy["month"] = pd.to_datetime(cdf_copy["contribution_date"]).dt.to_period("M").astype(str)
        contrib_monthly = (
            cdf_copy.groupby("month", as_index=False)
            .agg(contrib=("count", "sum"))
        )
        monthly = monthly.merge(contrib_monthly, on="month", how="left")
        monthly["contrib"] = monthly["contrib"].fillna(0)

    ws1.cell(row=1, column=1, value="Monthly Summary")
    ws1.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + (1 if have_contrib else 0))

    ws1.cell(row=2, column=1, value=f"Period: {start_date.strftime('%b %d, %Y')} — {end_date.strftime('%b %d, %Y')}")
    ws1.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

    headers = ["Month", "Tokens", "Cost ($)", "Requests"]
    if have_contrib:
        headers.append("Contributions")
    for col_idx, h in enumerate(headers, 1):
        ws1.cell(row=4, column=col_idx, value=h)
    _style_header(ws1, 4, len(headers))

    for row_idx, (_, row) in enumerate(monthly.iterrows(), 5):
        ws1.cell(row=row_idx, column=1, value=row["month"])
        _tok = ws1.cell(row=row_idx, column=2, value=int(row["tokens"]))
        _tok.number_format = TOKEN_NUMFMT
        ws1.cell(row=row_idx, column=3, value=round(row["cost"], 4))
        ws1.cell(row=row_idx, column=4, value=int(row["requests"]))
        if have_contrib:
            ws1.cell(row=row_idx, column=5, value=int(row["contrib"]))

    if len(monthly) > 0:
        _style_data_rows(ws1, 5, 4 + len(monthly), len(headers))
    _auto_width(ws1, len(headers))

    # --- Sheet 2: Per-User Allocation ---
    ws2 = wb.create_sheet("Per-User Allocation")

    user_agg = (
        df.groupby("username", as_index=False)
        .agg(tokens=("token_usage", "sum"), cost=("cost", "sum"), requests=("requests", "sum"))
        .sort_values("cost", ascending=False)
    )

    if have_contrib:
        contrib_by_user = (
            cdf.groupby("username", as_index=False)
            .agg(contrib=("count", "sum"))
        )
        user_agg = user_agg.merge(contrib_by_user, on="username", how="left")
        user_agg["contrib"] = user_agg["contrib"].fillna(0)

    ws2.cell(row=1, column=1, value="Per-User Cost Allocation")
    ws2.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5 + (1 if have_contrib else 0))

    ws2.cell(row=2, column=1, value=f"Period: {start_date.strftime('%b %d, %Y')} — {end_date.strftime('%b %d, %Y')}")
    ws2.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

    headers = ["S.No", "User", "Tokens", "Cost ($)", "Requests"]
    if have_contrib:
        headers.append("Contributions")
    for col_idx, h in enumerate(headers, 1):
        ws2.cell(row=4, column=col_idx, value=h)
    _style_header(ws2, 4, len(headers))

    for row_idx, (_, row) in enumerate(user_agg.iterrows(), 5):
        sno = row_idx - 4  # serial starting at 1
        ws2.cell(row=row_idx, column=1, value=sno)
        ws2.cell(row=row_idx, column=2, value=row["username"])
        _tok = ws2.cell(row=row_idx, column=3, value=int(row["tokens"]))
        _tok.number_format = TOKEN_NUMFMT
        ws2.cell(row=row_idx, column=4, value=round(row["cost"], 4))
        ws2.cell(row=row_idx, column=5, value=int(row["requests"]))
        if have_contrib:
            ws2.cell(row=row_idx, column=6, value=int(row["contrib"]))

    if len(user_agg) > 0:
        _style_data_rows(ws2, 5, 4 + len(user_agg), len(headers), left_cols=2)
    _auto_width(ws2, len(headers))

    # --- Sheet 3: Per-Model Breakdown ---
    ws3 = wb.create_sheet("Per-Model Breakdown")

    model_agg = (
        df.groupby("model_code", as_index=False)
        .agg(tokens=("token_usage", "sum"), cost=("cost", "sum"), requests=("requests", "sum"))
        .sort_values("cost", ascending=False)
    )

    ws3.cell(row=1, column=1, value="Per-Model Cost Breakdown")
    ws3.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    ws3.cell(row=2, column=1, value=f"Period: {start_date.strftime('%b %d, %Y')} — {end_date.strftime('%b %d, %Y')}")
    ws3.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

    headers = ["Model", "Tokens", "Cost ($)", "Requests"]
    for col_idx, h in enumerate(headers, 1):
        ws3.cell(row=4, column=col_idx, value=h)
    _style_header(ws3, 4, len(headers))

    for row_idx, (_, row) in enumerate(model_agg.iterrows(), 5):
        ws3.cell(row=row_idx, column=1, value=row["model_code"])
        _tok = ws3.cell(row=row_idx, column=2, value=int(row["tokens"]))
        _tok.number_format = TOKEN_NUMFMT
        ws3.cell(row=row_idx, column=3, value=round(row["cost"], 4))
        ws3.cell(row=row_idx, column=4, value=int(row["requests"]))

    if len(model_agg) > 0:
        _style_data_rows(ws3, 5, 4 + len(model_agg), len(headers))
    _auto_width(ws3, len(headers))

    # --- Sheet 4: Raw Data ---
    ws4 = wb.create_sheet("Raw Data")

    display_cols = [
        "billing_date", "username", "model_code", "token_type",
        "token_usage", "cost_price", "requests", "cost", "api_key", "billing_no",
    ]
    raw = df[display_cols].sort_values("billing_date")

    ws4.cell(row=1, column=1, value="Raw Data")
    ws4.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws4.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(display_cols))

    ws4.cell(row=2, column=1, value=f"Period: {start_date.strftime('%b %d, %Y')} — {end_date.strftime('%b %d, %Y')}")
    ws4.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

    headers = [
        "Date", "User", "Model", "Token Type",
        "Tokens", "Rate ($/kT)", "Requests", "Cost ($)", "API Key", "Billing No",
    ]
    for col_idx, h in enumerate(headers, 1):
        ws4.cell(row=4, column=col_idx, value=h)
    _style_header(ws4, 4, len(headers))

    for row_idx, (_, row) in enumerate(raw.iterrows(), 5):
        for col_idx, col_name in enumerate(display_cols, 1):
            val = row[col_name]
            if col_name in ("token_usage", "requests"):
                val = int(val)
            elif col_name in ("cost_price", "cost"):
                val = round(val, 4) if pd.notna(val) else 0
            _cell = ws4.cell(row=row_idx, column=col_idx, value=val)
            if col_name == "token_usage":
                _cell.number_format = TOKEN_NUMFMT

    if len(raw) > 0:
        _style_data_rows(ws4, 5, 4 + len(raw), len(headers))
    _auto_width(ws4, len(headers))

    # --- Sheet 5: Contributions (daily per-user detail) ---
    if have_contrib:
        ws5 = wb.create_sheet("Contributions")

        contrib_sorted = cdf.sort_values(["username", "contribution_date"])

        ws5.cell(row=1, column=1, value="Daily Contributions")
        ws5.cell(row=1, column=1).font = Font(bold=True, size=13)
        ws5.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

        ws5.cell(row=2, column=1, value=f"Period: {start_date.strftime('%b %d, %Y')} — {end_date.strftime('%b %d, %Y')}")
        ws5.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

        headers = ["User", "Date", "Contributions"]
        for col_idx, h in enumerate(headers, 1):
            ws5.cell(row=4, column=col_idx, value=h)
        _style_header(ws5, 4, len(headers))

        for row_idx, (_, row) in enumerate(contrib_sorted.iterrows(), 5):
            ws5.cell(row=row_idx, column=1, value=row["username"])
            ws5.cell(row=row_idx, column=2, value=str(row["contribution_date"]))
            ws5.cell(row=row_idx, column=3, value=int(row["count"]))

        if len(contrib_sorted) > 0:
            _style_data_rows(ws5, 5, 4 + len(contrib_sorted), len(headers))
        _auto_width(ws5, len(headers))

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
